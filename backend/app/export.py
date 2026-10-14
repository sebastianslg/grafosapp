"""Exportaciones: Excel (con fórmulas), Word (ecuaciones OMML), PDF, PNG/SVG, GraphML/GEXF, JSON, CSV, GeoJSON y
paquete de importación para Grafos v.1.3.5."""
from __future__ import annotations
import io
import json
import re
import zipfile
import networkx as nx
import numpy as np
from .algos.common import fmt
from .graphs import build_nx, undirected_view, separate_xy, auto_layout, edge_values
from .model import node_label, has_coords
import copy

MUTED = {"line": "7A7F87", "head": "E6E8EB"}


# ---------- Grafos v.1.3.5 ----------

def _crlf(rows):
    return ("\r\n".join("\t".join(r) for r in rows) + "\r\n").encode("utf-8")


def grafos_package(project: dict, scale: int | None = None, min_sep: bool = True, criterion: str = "cost",
                   value_field: str = "demand", source: str | None = None, target: str | None = None) -> dict[str, bytes]:
    """Archivos de importación para Grafos v.1.3.5 + instrucciones con valores esperados (calculados aquí)."""
    p = copy.deepcopy(project)
    sc = int(scale or p["params"].get("scale") or 1)
    ids = [n["id"] for n in p["nodes"]]
    idx = {v: i for i, v in enumerate(ids)}
    n = len(ids)
    g = build_nx(p, criterion)
    M = [["#"] * n for _ in range(n)]
    W = {}
    for u, v, a in g.edges(data=True):
        w = int(round(a["w"] * sc))
        M[idx[u]][idx[v]] = str(w)
        W[(idx[u], idx[v])] = w
    CAP = [["#"] * n for _ in range(n)]
    ncap = 0
    for u, v, a in g.edges(data=True):
        if a.get("capacity") is not None:
            CAP[idx[u]][idx[v]] = str(int(round(a["capacity"])))
            ncap += 1
    files = {"1_matriz_coste_arcos.txt": _crlf(M)}
    if ncap:
        files["2_matriz_maximo_arcos.txt"] = _crlf(CAP)
    auto_layout(p)
    if min_sep and all(x.get("x") is not None for x in p["nodes"]):
        separate_xy(p, 10.0)
    if all(x.get("x") is not None for x in p["nodes"]):
        files["3_coordenadas_XY_nodos.txt"] = _crlf([[str(i), str(int(round(x["x"]))), str(int(round(x["y"])))] for i, x in enumerate(p["nodes"])])
    vals = [(i, nd.get(value_field)) for i, nd in enumerate(p["nodes"]) if nd.get(value_field) is not None]
    if vals:
        files["4_valores_nodo_2col.txt"] = _crlf([[str(i), str(int(round(v)))] for i, v in vals])
    # valores de control sobre los enteros exportados (lo que Grafos realmente leerá)
    G = nx.DiGraph()
    G.add_nodes_from(range(n))
    for (a, b), w in W.items():
        G.add_edge(a, b, w=w)
    ctrl = []
    sp = dict(nx.all_pairs_dijkstra_path_length(G, weight="w"))
    tot = sum(d for a in sp for b, d in sp[a].items() if a != b)
    ctrl.append(("Suma de todos los pares ordenados (Floyd-Warshall, pares conectados)", tot))
    s = idx.get(source, 0) if source else 0
    d0 = sp[s]
    ctrl.append((f"Costos mínimos desde el nodo {s} ({node_label(p, ids[s])}) a todos", [d0.get(j) for j in range(n)]))
    if target and target in idx:
        t = idx[target]
        if t in d0:
            pth = nx.dijkstra_path(G, s, t, weight="w")
            ctrl.append((f"Dijkstra {s}→{t}: costo", d0[t]))
            ctrl.append((f"Dijkstra {s}→{t}: camino", " > ".join(map(str, pth))))
    U = nx.Graph()
    for (a, b), w in W.items():
        if U.has_edge(a, b):
            U[a][b]["w"] = min(U[a][b]["w"], w)
        else:
            U.add_edge(a, b, w=w)
    if U.number_of_edges():
        T = nx.minimum_spanning_tree(U, weight="w")
        ctrl.append((f"Árbol de valor total mínimo (suma de aristas, cada arista una vez; {T.number_of_edges()} aristas)",
                     sum(d["w"] for _, _, d in T.edges(data=True))))
        T2 = nx.maximum_spanning_tree(U, weight="w")
        ctrl.append(("Árbol de valor total máximo (cada arista una vez)", sum(d["w"] for _, _, d in T2.edges(data=True))))
    txt = ["INSTRUCCIONES DE CARGA EN GRAFOS v.1.3.5",
           "=" * 44, "",
           f"Proyecto: {p['name']}  |  Nodos: {n}  |  Arcos en la matriz: {len(W)}  |  Factor de escala: ×{sc} (enteros)",
           f"Criterio de la matriz de coste: {criterion}",
           "Formato: separador TAB, saltos de línea CRLF, '#' = sin valor.",
           "Nota: el índice de cada nodo en Grafos es el de esta tabla (0 = primer nodo).", "",
           "TABLA DE NODOS (índice → id → nombre)"]
    txt += [f"  {i}\t{ids[i]}\t{p['nodes'][i].get('name', '')}" for i in range(n)]
    txt += ["", "ORDEN DE CARGA (una carga a la vez, por separado)", "",
            "1) Matriz coste (arcos) — archivo 1_matriz_coste_arcos.txt",
            "   Menú: importar → «Matriz coste (arcos)» · Modo: SUSTITUIR."]
    if ncap:
        txt += ["2) Matriz máximo (arcos) — archivo 2_matriz_maximo_arcos.txt", "   Menú: importar → «Matriz máximo (arcos)» · Modo: SUSTITUIR. (capacidades para flujo máximo)"]
    if "3_coordenadas_XY_nodos.txt" in files:
        txt += ["3) Coordenadas — archivo 3_coordenadas_XY_nodos.txt (i<TAB>x<TAB>y)",
                "   Menú: importar → «Matriz etiqueta y coord XY (nodos)» · Modo: ACTUALIZAR (con SUSTITUIR se pierden los arcos).",
                "   Los nodos quedan separados al menos 10 px." if min_sep else ""]
    if "4_valores_nodo_2col.txt" in files:
        txt += [f"4) Valores de nodo ({value_field}) — archivo 4_valores_nodo_2col.txt (i<TAB>valor, 2 columnas)",
                "   Menú: importar → «Matriz etiqueta y valor (nodos)» · Modo: ACTUALIZAR."]
    txt += ["", "USO", "  Clic IZQUIERDO sobre un nodo = ORIGEN (amarillo). Clic DERECHO = DESTINO (rojo).",
            "  Caminos: menú de caminos → Camino mínimo – Alg. Dijkstra (origen–destino); la ventana lista el costo a todos los nodos desde el origen.",
            "  Todos los caminos mínimos – Alg. FloydWarshall; árboles: Árbol de valor total mínimo/máximo – Alg. Kruskal o Prim.",
            "  Flujo máximo – Alg. FordFulkerson (origen–destino) usa la matriz máximo.", "",
            "VALORES ESPERADOS (calculados sobre los enteros exportados)"]
    for name, v in ctrl:
        txt.append(f"  • {name}: {v}")
    txt += ["", "NO VERIFICADO: la forma exacta en que Grafos presenta el valor del árbol (una vez por arista vs. dos arcos) y el CVRP/VRP-XML",
            "no se ha comprobado contra la interfaz; el CVRP requiere crear la estructura VRP-XML en «Gestión de datos VRP-XML»."]
    files["0_INSTRUCCIONES_Grafos.txt"] = ("\r\n".join(t for t in txt if t is not None)).encode("utf-8")
    return files


# ---------- Excel ----------

def to_excel(project: dict, results: list[dict] | None = None, scenarios: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    wb = Workbook()
    thin = Side(style="thin", color=MUTED["line"])
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    hf = PatternFill("solid", fgColor=MUTED["head"])

    def sheet(title, header, rows):
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.border = Font(bold=True), hf, bd
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = bd
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or "")) for c in col[:60]) + 2))
        ws.freeze_panes = "A2"
        return ws
    wb.remove(wb.active)
    pr = project["params"]
    wsP = sheet("Parámetros", ["Parámetro", "Valor", "Descripción"],
                [["c_km", pr.get("c_km"), "Costo por km"], ["c_hora", pr.get("c_hora"), "Costo por hora"], ["speed", pr.get("speed"), "Velocidad km/h"],
                 ["road_factor", pr.get("road_factor"), "Factor vial"], ["e_km", pr.get("e_km"), "Emisiones por km"], ["scale", pr.get("scale"), "Factor de escala (enteros)"],
                 ["depots", ", ".join(pr.get("depots", [])), "Depósitos"]])
    sheet("Nodos", ["id", "nombre", "lat", "lon", "x", "y", "demanda", "oferta", "tipo", "tw_inicio", "tw_fin", "servicio", "capacidad", "frecuencia"],
          [[n["id"], n.get("name"), n.get("lat"), n.get("lon"), n.get("x"), n.get("y"), n.get("demand"), n.get("supply"), n.get("type"),
            n.get("tw_start"), n.get("tw_end"), n.get("service"), n.get("capacity"), n.get("frequency")] for n in project["nodes"]])
    # Aristas con fórmulas visibles (tiempo y costo derivados si hay distancia)
    by = {n["id"]: n for n in project["nodes"]}
    rows = []
    for i, e in enumerate(project["edges"], start=2):
        v = edge_values(project, e, by)
        d_cell = v["distance"]
        t = e.get("time") if e.get("time") is not None else (f"=60*D{i}/Parámetros!$B$4" if d_cell is not None else None)
        c = e.get("cost") if e.get("cost") is not None else (f"=Parámetros!$B$2*D{i}+Parámetros!$B$3*E{i}/60" if d_cell is not None else None)
        em = e.get("emissions") if e.get("emissions") is not None else (f"=Parámetros!$B$6*D{i}" if d_cell is not None else None)
        rows.append([e["source"], e["target"], "sí" if e.get("directed") else "no", d_cell, t, c, e.get("capacity"), em])
    sheet("Aristas", ["origen", "destino", "dirigida", "distancia_km", "tiempo_min", "costo", "capacidad", "emisiones"], rows)
    if project["vehicles"]:
        sheet("Vehículos", ["tipo", "capacidad", "costo_fijo", "costo_km", "velocidad", "autonomía", "cantidad", "emisiones_km", "depósito"],
              [[v.get("type"), v.get("capacity"), v.get("fixed_cost"), v.get("cost_per_km"), v.get("speed"), v.get("range"), v.get("count"),
                v.get("emissions_per_km"), v.get("depot")] for v in project["vehicles"]])
    try:
        g = build_nx(project, "cost")
        ids = list(g.nodes)
        D = dict(nx.all_pairs_dijkstra_path_length(g, weight="w"))
        if len(ids) <= 300:
            sheet("Matriz de costos", ["↓ origen / destino →"] + ids, [[a] + [round(D[a].get(b), 6) if b in D[a] else None for b in ids] for a in ids])
    except Exception:
        pass
    for k, r in enumerate(results or [], 1):
        base = f"R{k}_{r.get('key', r['algorithm'])}"[:25]
        sheet(base, ["Campo", "Valor"], [["Algoritmo", r["title"]], ["Estado", r["status"]], ["Valor", r.get("value")], ["Mensaje", r.get("message")],
                                         ["Parámetros", json.dumps(r.get("run_params", {}), ensure_ascii=False, default=str)]] +
              [["Control: " + c["name"], json.dumps(c["value"], default=str)] for c in r.get("control", [])])
        for ti, t in enumerate(r.get("tables", [])[:8]):
            sheet(f"{base}_T{ti + 1}"[:31], t["columns"], t["rows"])
        if r.get("matrix"):
            m = r["matrix"]
            sheet(f"{base}_Mat"[:31], [""] + m["labels"], [[m["labels"][i]] + row for i, row in enumerate(m["values"])])
    if scenarios:
        sheet("Escenarios", ["Escenario", "Valor", "Δ %", "Algoritmo", "Estado", "Mensaje"],
              [[r["escenario"], r["valor"], r["delta_pct"], r["algoritmo"], r["estado"], r["mensaje"]] for r in scenarios["rows"]])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# ---------- imagen del grafo ----------

def _positions(project: dict):
    ns = project["nodes"]
    if all(n.get("lon") is not None and n.get("lat") is not None for n in ns):
        return {n["id"]: (n["lon"], n["lat"]) for n in ns}, "geo"
    q = copy.deepcopy(project)
    auto_layout(q)
    return {n["id"]: (n["x"], -n["y"]) for n in q["nodes"]}, "xy"


def graph_image(project: dict, result: dict | None = None, fmt_: str = "png", labels: bool = True) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    pos, _ = _positions(project)
    fig, ax = plt.subplots(figsize=(9, 6.5), dpi=140)
    for e in project["edges"]:
        if e["source"] in pos and e["target"] in pos:
            (x1, y1), (x2, y2) = pos[e["source"]], pos[e["target"]]
            ax.plot([x1, x2], [y1, y2], color="#C3C7CD", lw=0.8, zorder=1)
    palette = ["#8C2D2D", "#2D5F8C", "#3F7D4E", "#B0752B", "#6B4E8C", "#2B7F83", "#8A6D1F", "#5A5A5A"]
    hn = set()
    if result:
        groups: dict[str, int] = {}
        for a, b, grp in result.get("highlight_edges", []):
            if a in pos and b in pos:
                k = groups.setdefault(grp, len(groups))
                (x1, y1), (x2, y2) = pos[a], pos[b]
                ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="-|>", color=palette[k % 8], lw=2, shrinkA=4, shrinkB=4), zorder=3)
        hn = set(result.get("highlight_nodes", []))
        for gname, k in list(groups.items())[:8]:
            ax.plot([], [], color=palette[k % 8], lw=2, label=gname)
        if groups:
            ax.legend(loc="best", fontsize=7, frameon=True)
    dep = set(project["params"].get("depots", []))
    for n in project["nodes"]:
        x, y = pos[n["id"]]
        ax.scatter([x], [y], s=70 if n["id"] in dep else 38, c="#222" if n["id"] in dep else ("#D9A441" if n["id"] in hn else "#6D7B8D"),
                   marker="s" if n["id"] in dep else "o", zorder=4, edgecolors="white", linewidths=0.6)
        if labels:
            ax.annotate(str(n["id"]), (x, y), xytext=(4, 4), textcoords="offset points", fontsize=6.5, zorder=5)
    ax.set_axis_off()
    ax.set_title((result or {}).get("title", project["name"]), color="black", fontsize=11)
    b = io.BytesIO()
    fig.savefig(b, format=fmt_, bbox_inches="tight")
    plt.close(fig)
    return b.getvalue()


# ---------- Word ----------

def _omml(text: str):
    """Ecuación lineal → OMML (subíndices/superíndices con _{...} y ^{...})."""
    from lxml import etree
    M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    nsm = {"m": M}
    root = etree.Element(f"{{{M}}}oMathPara", nsmap=nsm)
    om = etree.SubElement(root, f"{{{M}}}oMath")

    def run(parent, s):
        if not s:
            return
        r = etree.SubElement(parent, f"{{{M}}}r")
        t = etree.SubElement(r, f"{{{M}}}t")
        t.text = s
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    pos = 0
    pat = re.compile(r"(\S)([_^])\{([^}]*)\}")
    for m in pat.finditer(text):
        run(om, text[pos:m.start()])
        tag = "sSub" if m.group(2) == "_" else "sSup"
        el = etree.SubElement(om, f"{{{M}}}{tag}")
        e = etree.SubElement(el, f"{{{M}}}e")
        run(e, m.group(1))
        s = etree.SubElement(el, f"{{{M}}}{'sub' if m.group(2) == '_' else 'sup'}")
        run(s, m.group(3))
        pos = m.end()
    run(om, text[pos:])
    return root


def _set_borders(table, color="7A7F87"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def _shade(cell, fill="E6E8EB"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def to_word(project: dict, result: dict, explanation: dict, image: bytes | None = None, template: bytes | None = None) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = docx.Document(io.BytesIO(template)) if template else docx.Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for lv in (1, 2, 3):
        try:
            h = d.styles[f"Heading {lv}"]
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.font.name = "Calibri"
        except KeyError:
            pass
    t = d.add_heading(result["title"], 0)
    for r in t.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
    d.add_paragraph(f"Proyecto: {project['name']} · {len(project['nodes'])} nodos · {len(project['edges'])} aristas")
    d.add_heading("Resumen ejecutivo", 1)
    for ln in explanation["executive"]["short"]:
        d.add_paragraph(ln, style="List Bullet")
    d.add_heading("Respuesta directa", 1)
    d.add_paragraph(explanation["answer"])
    if image:
        d.add_picture(io.BytesIO(image), width=Inches(6.2))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading("Paso a paso entre nodos", 1)
    for blk in explanation["steps"][:10]:
        d.add_heading(blk["label"], 3)
        d.add_paragraph(blk["summary"])
        for ln in blk["lines"][:60]:
            d.add_paragraph(ln, style="List Number")
    d.add_heading("Razonamiento de elección", 1)
    r = explanation["reasoning"]
    d.add_paragraph(r["algorithm_choice"])
    if r["equation"]:
        p = d.add_paragraph()
        p._p.append(_omml(r["equation"]))
    for k in ("optimality", "comparison"):
        if r.get(k):
            d.add_paragraph(r[k])
    d.add_paragraph("Restricciones activas: " + "; ".join(r["active_constraints"]))
    d.add_heading("Informe técnico", 1)
    for sec in explanation["technical"]:
        d.add_heading(sec["heading"], 2)
        for ptxt in sec["paragraphs"]:
            if ptxt:
                d.add_paragraph(ptxt)
        for eq in sec.get("equations", []):
            p = d.add_paragraph()
            p._p.append(_omml(eq))
        for tb in sec.get("tables", []):
            _add_table(d, tb)
    d.add_heading("Valores de control", 1)
    tbl = d.add_table(rows=1, cols=2)
    _set_borders(tbl)
    tbl.rows[0].cells[0].text, tbl.rows[0].cells[1].text = "Indicador", "Valor"
    for c in tbl.rows[0].cells:
        _shade(c)
        c.paragraphs[0].runs[0].font.bold = True
    for c in explanation["control"]:
        row = tbl.add_row().cells
        row[0].text, row[1].text = c["name"], json.dumps(c["value"], default=str, ensure_ascii=False)[:300]
    d.add_paragraph()
    d.add_paragraph(f"Reproducibilidad: parámetros {json.dumps(result.get('run_params', {}), ensure_ascii=False, default=str)[:500]}")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _add_table(d, tb, max_rows=80):
    from docx.shared import Pt
    d.add_paragraph(tb["title"]).runs[0].font.bold = True
    t = d.add_table(rows=1, cols=len(tb["columns"]))
    _set_borders(t)
    for i, c in enumerate(tb["columns"]):
        t.rows[0].cells[i].text = str(c)
        _shade(t.rows[0].cells[i])
        t.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
    for row in tb["rows"][:max_rows]:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for rn in p.runs:
                    rn.font.size = Pt(9)
    if len(tb["rows"]) > max_rows:
        d.add_paragraph(f"(… {len(tb['rows']) - max_rows} filas más en el Excel)")


# ---------- PDF ----------

def to_pdf(project: dict, result: dict, explanation: dict, image: bytes | None = None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
    from xml.sax.saxutils import escape
    ss = getSampleStyleSheet()
    H = ParagraphStyle("H", parent=ss["Heading2"], textColor=colors.black)
    T = ParagraphStyle("T", parent=ss["Title"], textColor=colors.black)
    b = io.BytesIO()
    doc = SimpleDocTemplate(b, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm, title=result["title"])
    S = [Paragraph(escape(result["title"]), T), Paragraph(escape(f"Proyecto: {project['name']}"), ss["Normal"]), Spacer(1, 8),
         Paragraph("Resumen ejecutivo", H)]
    S += [Paragraph("• " + escape(x), ss["Normal"]) for x in explanation["executive"]["short"]]
    S += [Paragraph("Respuesta directa", H), Paragraph(escape(explanation["answer"]), ss["Normal"])]
    if image:
        S.append(Image(io.BytesIO(image), width=15 * cm, height=10.5 * cm))
    S.append(Paragraph("Paso a paso entre nodos", H))
    for blk in explanation["steps"][:6]:
        S.append(Paragraph(f"<b>{escape(blk['label'])}</b>: {escape(blk['summary'][:300])}", ss["Normal"]))
        S += [Paragraph(f"{i}. {escape(l)}", ss["Normal"]) for i, l in enumerate(blk["lines"][:40], 1)]
    r = explanation["reasoning"]
    S += [Paragraph("Razonamiento de elección", H), Paragraph(escape(r["algorithm_choice"]), ss["Normal"])]
    if r["equation"]:
        S.append(Paragraph("<i>" + escape(r["equation"]) + "</i>", ss["Normal"]))
    for k in ("optimality", "comparison"):
        if r.get(k):
            S.append(Paragraph(escape(r[k]), ss["Normal"]))
    for tb in result.get("tables", [])[:5]:
        S.append(Paragraph(escape(tb["title"]), H))
        data = [tb["columns"]] + [[str(c)[:40] for c in row] for row in tb["rows"][:40]]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#7A7F87")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6E8EB")),
                               ("FONTSIZE", (0, 0), (-1, -1), 7.5)]))
        S.append(t)
    S.append(Paragraph("Valores de control", H))
    data = [["Indicador", "Valor"]] + [[c["name"], json.dumps(c["value"], default=str, ensure_ascii=False)[:120]] for c in explanation["control"]]
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#7A7F87")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6E8EB")), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    S.append(t)
    doc.build(S)
    return b.getvalue()


# ---------- formatos de grafo / tabulares ----------

def to_graphml(project: dict, fmt_: str = "graphml") -> bytes:
    g = nx.DiGraph() if any(e.get("directed") for e in project["edges"]) else nx.Graph()
    for n in project["nodes"]:
        g.add_node(n["id"], **{k: v for k, v in n.items() if k not in ("id", "extra") and v is not None})
    by = {n["id"]: n for n in project["nodes"]}
    for e in project["edges"]:
        v = edge_values(project, e, by)
        g.add_edge(e["source"], e["target"], **{k: float(x) for k, x in v.items() if x is not None},
                   **({"capacity": float(e["capacity"])} if e.get("capacity") is not None else {}))
    b = io.BytesIO()
    (nx.write_graphml if fmt_ == "graphml" else nx.write_gexf)(g, b)
    return b.getvalue()


def to_geojson(project: dict, result: dict | None = None) -> bytes:
    if not has_coords(project):
        raise ValueError("GeoJSON requiere lat/lon en todos los nodos.")
    feats = [{"type": "Feature", "geometry": {"type": "Point", "coordinates": [n["lon"], n["lat"]]},
              "properties": {"id": n["id"], "name": n.get("name"), "type": n.get("type"), "demand": n.get("demand")}} for n in project["nodes"]]
    by = {n["id"]: n for n in project["nodes"]}
    for e in project["edges"]:
        a, b = by.get(e["source"]), by.get(e["target"])
        if a and b:
            feats.append({"type": "Feature", "geometry": {"type": "LineString", "coordinates": [[a["lon"], a["lat"]], [b["lon"], b["lat"]]]},
                          "properties": {"source": e["source"], "target": e["target"], "distance": e.get("distance"), "kind": "arista"}})
    if result:
        for pth in result.get("paths", []):
            co = [[by[x]["lon"], by[x]["lat"]] for x in pth["nodes"] if x in by]
            if len(co) > 1:
                feats.append({"type": "Feature", "geometry": {"type": "LineString", "coordinates": co},
                              "properties": {"kind": "ruta", "label": pth.get("label"), "cost": pth.get("cost")}})
    return json.dumps({"type": "FeatureCollection", "features": feats}, ensure_ascii=False).encode()


def to_csv_zip(project: dict) -> bytes:
    import pandas as pd
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        nodes = [{k: v for k, v in n.items() if k != "extra"} for n in project["nodes"]]
        z.writestr("nodos.csv", pd.DataFrame(nodes).to_csv(index=False))
        z.writestr("aristas.csv", pd.DataFrame(project["edges"]).to_csv(index=False))
        z.writestr("vehiculos.csv", pd.DataFrame(project["vehicles"]).to_csv(index=False))
    return b.getvalue()


def to_sqlite(project: dict) -> bytes:
    import sqlite3, tempfile, os
    import pandas as pd
    with tempfile.NamedTemporaryFile(delete=False, suffix=".db") as f:
        path = f.name
    try:
        con = sqlite3.connect(path)
        pd.DataFrame([{k: v for k, v in n.items() if k != "extra"} for n in project["nodes"]]).to_sql("nodos", con, index=False)
        pd.DataFrame(project["edges"]).to_sql("aristas", con, index=False)
        pd.DataFrame(project["vehicles"] or [{"tipo": None}]).to_sql("vehiculos", con, index=False)
        pd.DataFrame([{"clave": k, "valor": json.dumps(v)} for k, v in project["params"].items()]).to_sql("parametros", con, index=False)
        con.close()
        return open(path, "rb").read()
    finally:
        os.unlink(path)


TEMPLATES = {
    "nodos_aristas": {"nodos.csv": "id,nombre,lat,lon,demanda,oferta,tipo,tw_inicio,tw_fin,servicio,frecuencia\nA,Depósito,10.40,-75.50,0,,depósito,,,0,\nB,Cliente B,10.42,-75.48,50,,cliente,0,480,5,1\nC,Cliente C,10.45,-75.52,70,,cliente,0,480,5,2\n",
                      "aristas.csv": "origen,destino,distancia_km,tiempo_min,costo,capacidad,dirigido\nA,B,4.2,,,60,no\nB,C,6.1,,,40,no\nA,C,8.3,,,50,no\n"},
    "matriz_costos": {"matriz.csv": "origen,A,B,C\nA,0,4.2,8.3\nB,4.2,0,6.1\nC,8.3,6.1,0\n"},
    "flota": {"vehiculos.csv": "tipo,capacidad,costo_fijo,costo_km,velocidad,autonomia,cantidad,emisiones_km,deposito\nFurgón diésel,500,30,1.2,40,,3,0.22,A\nFurgoneta eléctrica,300,20,0.6,35,150,2,0,A\n"},
    "transbordo": {"nodos.csv": "id,oferta,demanda\nS1,70,\nS2,50,\nT1,,\nD1,,60\nD2,,60\n",
                   "aristas.csv": "origen,destino,costo,capacidad,dirigido\nS1,T1,4,80,si\nS2,T1,5,50,si\nT1,D1,2,60,si\nT1,D2,3,60,si\n"},
    "parametros": {"parametros.csv": "parametro,valor\ncosto por km,1.2\ncosto por hora,6\nvelocidad,40\nfactor vial,1.25\n"},
}


def template_zip(kind: str) -> bytes:
    if kind not in TEMPLATES:
        raise ValueError(f"Plantilla desconocida: {kind}. Disponibles: {list(TEMPLATES)}")
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for n, c in TEMPLATES[kind].items():
            z.writestr(n, c)
    return b.getvalue()


def full_zip(project: dict, results: list[dict], explanations: list[dict], scenarios: dict | None = None, **grafos_kw) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("proyecto.json", json.dumps(project, ensure_ascii=False, indent=1, default=str))
        z.writestr("datos_canonicos.xlsx", to_excel(project, results, scenarios))
        z.writestr("grafo.graphml", to_graphml(project))
        z.writestr("grafo.png", graph_image(project, results[-1] if results else None))
        for name, data in grafos_package(project, **grafos_kw).items():
            z.writestr(f"Grafos_v1.3.5/{name}", data)
        for i, (r, e) in enumerate(zip(results, explanations), 1):
            img = graph_image(project, r)
            z.writestr(f"informe_{i}_{r.get('key', 'x')}.docx", to_word(project, r, e, img))
            z.writestr(f"informe_{i}_{r.get('key', 'x')}.pdf", to_pdf(project, r, e, img))
            z.writestr(f"resultado_{i}_{r.get('key', 'x')}.json", json.dumps(r, ensure_ascii=False, default=str, indent=1))
        if has_coords(project):
            z.writestr("grafo.geojson", to_geojson(project, results[-1] if results else None))
    return b.getvalue()
