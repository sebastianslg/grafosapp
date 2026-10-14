import io, json, sqlite3, zipfile
import pytest
from app import ingest
from app.ingest import parse_file, build_project, preview, IngestError
from app.model import validate


def build(*files, **kw):
    return build_project("t", [parse_file(n, r) for n, r in files], **kw)


def test_grafos_files_roundtrip():
    mat = b"#\t5\t3\r\n5\t#\t2\r\n3\t2\t#\r\n"
    xy = b"0\t100\t100\r\n1\t200\t150\r\n2\t120\t300\r\n"
    val = b"0\t10\r\n1\t20\r\n2\t30\r\n"
    p = build(("m.txt", mat), ("xy.txt", xy), ("v.txt", val))
    assert len(p["nodes"]) == 3 and len(p["edges"]) == 3
    assert p["nodes"][1]["x"] == 200 and p["nodes"][1]["y"] == 150
    assert p["nodes"][2]["demand"] == 30
    assert {(e["source"], e["target"]) for e in p["edges"]} == {("0", "1"), ("0", "2"), ("1", "2")}
    assert all(e["directed"] is False for e in p["edges"])


def test_asymmetric_matrix_is_directed():
    p = build(("m.txt", b"#\t4\t#\n#\t#\t2\n1\t#\t#\n"))
    assert len(p["edges"]) == 3 and all(e["directed"] for e in p["edges"])


def test_csv_spanish_headers_semicolon_latin1_decimal_comma():
    raw = "Código;Nombre;Latitud;Longitud;Demanda (kg)\nA;Bogotá;4,711;-74,07;0\nB;Medellín;6,25;-75,56;50\nC;Cali;3,45;-76,53;70\n".encode("latin-1")
    pv = preview([parse_file("sitios.csv", raw)])
    t = pv["tables"][0]
    assert t["role"] == "nodes" and t["mapping"]["lat"] == "Latitud" and t["mapping"]["demand"] == "Demanda (kg)"
    p = build(("sitios.csv", raw))
    assert p["nodes"][0]["lat"] == pytest.approx(4.711) and p["nodes"][1]["name"] == "Medellín"


def test_edges_table_without_node_table_creates_nodes():
    raw = b"from,to,km,time_min\nA,B,3.5,6\nB,C,2.0,4\n"
    p = build(("e.csv", raw))
    assert [n["id"] for n in p["nodes"]] == ["A", "B", "C"] and p["edges"][0]["distance"] == 3.5 and p["edges"][0]["time"] == 6


def test_headerless_edge_list():
    p = build(("g.edgelist", b"a b 3\nb c 4\n"))
    assert len(p["edges"]) == 2


def test_unit_conversion_meters():
    p = build(("e.csv", b"origen,destino,distancia\nA,B,2500\n"), units={"distance_divisor": 1000})
    assert p["edges"][0]["distance"] == 2.5


def test_excel_multisheet_and_fleet():
    import pandas as pd
    b = io.BytesIO()
    with pd.ExcelWriter(b) as w:
        pd.DataFrame({"id": ["D", "X", "Y"], "lat": [1, 1.1, 1.2], "lon": [2, 2.1, 2.3], "demanda": [0, 5, 6]}).to_excel(w, sheet_name="Sitios", index=False)
        pd.DataFrame({"tipo": ["Van"], "capacidad": [100], "costo km": [1.5], "cantidad": [2]}).to_excel(w, sheet_name="Flota", index=False)
    p = build(("datos.xlsx", b.getvalue()))
    assert len(p["nodes"]) == 3 and p["vehicles"][0]["capacity"] == 100 and p["vehicles"][0]["count"] == 2


def test_json_geojson_graphml():
    js = json.dumps({"nodes": [{"id": 1, "lat": 4, "lon": -74}, {"id": 2, "lat": 4.1, "lon": -74.1}], "edges": [{"source": 1, "target": 2, "distance": 3}]}).encode()
    assert len(build(("g.json", js))["edges"]) == 1
    gj = json.dumps({"type": "FeatureCollection", "features": [
        {"type": "Feature", "properties": {"id": "P"}, "geometry": {"type": "Point", "coordinates": [-74, 4]}},
        {"type": "Feature", "properties": {"id": "Q"}, "geometry": {"type": "Point", "coordinates": [-74.1, 4.1]}}]}).encode()
    assert len(build(("p.geojson", gj))["nodes"]) == 2
    from app.samples import sample_small
    from app import export
    gm = export.to_graphml(sample_small())
    p = build(("g.graphml", gm))
    assert len(p["nodes"]) == 14


def test_sqlite_and_sql_dump_rejects_unsafe_statements():
    con = sqlite3.connect(":memory:")
    con.execute("create table nodos(id text, lat real, lon real)")
    con.executemany("insert into nodos values(?,?,?)", [("a", 1, 2), ("b", 1.1, 2.1)])
    import tempfile, os
    f = tempfile.mktemp()
    d = sqlite3.connect(f)
    d.execute("create table nodos(id text, lat real, lon real)"); d.execute("insert into nodos values('a',1,2),('b',1.1,2.1)"); d.commit(); d.close()
    assert len(build(("x.db", open(f, "rb").read()))["nodes"]) == 2
    os.unlink(f)
    dump = b"CREATE TABLE nodos (id TEXT, lat REAL, lon REAL);\nINSERT INTO nodos VALUES ('a',1,2);\nINSERT INTO nodos VALUES ('b',1.1,2.1);\nDROP TABLE nodos;\nATTACH DATABASE '/tmp/x' AS y;\n"
    ps = parse_file("dump.sql", dump)
    assert any("omitidas" in n for n in ps.notes)
    assert len(build_project("t", [ps])["nodes"]) == 2


@pytest.mark.parametrize("name,raw", [("a.csv", b""), ("b.xlsx", b"PK\x03\x04garbage"), ("c.json", b"{not json"), ("d.pdf", b"%PDF-1.4 roto"),
                                       ("e.db", b"SQLite format 3\x00" + b"\x00" * 50), ("f.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1xxxx"), ("g.graphml", b"<graphml>")])
def test_malformed_files_raise_clear_errors(name, raw):
    with pytest.raises(IngestError):
        parse_file(name, raw)


def test_too_large_file(monkeypatch):
    monkeypatch.setattr(ingest, "MAX_BYTES", 100)
    with pytest.raises(IngestError, match="grande"):
        parse_file("x.csv", b"a,b\n" * 200)


def test_validation_warnings():
    p = build(("e.csv", b"origen,destino,costo\nA,B,-3\n"))
    codes = {w["codigo"] for w in validate(p)}
    assert "valores_negativos" in codes
    p["edges"].append({"source": "A", "target": "ZZ", "distance": 1, "time": None, "cost": None, "capacity": None, "directed": False, "emissions": None})
    assert "aristas_huerfanas" in {w["codigo"] for w in validate(p)}
    p["nodes"][0]["lat"], p["nodes"][0]["lon"] = 4.5, -200
    assert "coordenadas_fuera_de_rango" in {w["codigo"] for w in validate(p)}


def test_docx_report_extracts_tables_and_params_with_citations():
    import docx
    d = docx.Document()
    d.add_paragraph("Costo por km: 1.35")
    d.add_paragraph("Velocidad media: 38")
    d.add_paragraph("Texto sin datos relevantes.")
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["id", "lat", "lon"], ["P1", "10.1", "-75.1"], ["P2", "10.2", "-75.2"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    b = io.BytesIO(); d.save(b)
    ps = parse_file("informe.docx", b.getvalue())
    pv = preview([ps])
    params = {x["param"]: x for x in pv["extracted_params"]}
    assert params["c_km"]["valor"] == 1.35 and "párrafo 1" in params["c_km"]["cita"] and params["speed"]["valor"] == 38
    assert "road_factor" not in params                     # lo ausente no se inventa
    p = build_project("t", [ps], extra_params={k: v["valor"] for k, v in params.items()})
    assert len(p["nodes"]) == 2 and p["params"]["c_km"] == 1.35


def test_pdf_report_tables_and_text():
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table
    from reportlab.lib.styles import getSampleStyleSheet
    b = io.BytesIO()
    SimpleDocTemplate(b).build([Paragraph("Costo por hora: 12.5", getSampleStyleSheet()["Normal"]),
                                Table([["id", "lat", "lon"], ["A", "1", "2"], ["B", "1.5", "2.5"]], style=[("GRID", (0, 0), (-1, -1), .5, "black")])])
    ps = parse_file("r.pdf", b.getvalue())
    assert any(x["param"] == "c_hora" and x["valor"] == 12.5 for x in preview([ps])["extracted_params"])
    assert ps.tables and len(ps.tables[0].df) == 2


def test_matrix_with_labels():
    raw = b"X,A,B,C\nA,0,1,2\nB,1,0,3\nC,2,3,0\n"
    p = build(("m.csv", raw))
    assert len(p["nodes"]) == 3 and len(p["edges"]) == 3
