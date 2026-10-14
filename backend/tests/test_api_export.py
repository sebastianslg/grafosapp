import io, json, zipfile
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app import export
from app.samples import sample_small

c = TestClient(app)


@pytest.fixture(scope="module")
def pid():
    return c.post("/api/samples/pequeno").json()["id"]


def test_catalog_has_all_families():
    cats = {a["category"] for a in c.get("/api/catalog").json()}
    assert {"Caminos", "Árboles", "Flujos", "Estructura", "Localización", "Rutas", "Escenarios"} <= cats
    assert len(c.get("/api/catalog").json()) >= 40


def test_run_persist_and_reproduce(pid):
    body = {"algorithm": "tsp", "params": {"depot": "0", "time_limit": 2, "seed": 5}}
    a = c.post(f"/api/projects/{pid}/run?wait=true", json=body).json()
    b = c.post(f"/api/projects/{pid}/run?wait=true", json=body).json()
    assert a["result"]["value"] == b["result"]["value"]
    run = c.get(f"/api/runs/{a['run_id']}").json()
    assert run["params"]["seed"] == 5 and run["result"]["explanation"]["answer"]


def test_async_job_progress_and_cancel(pid):
    j = c.post(f"/api/projects/{pid}/run", json={"algorithm": "kruskal", "params": {}}).json()["job_id"]
    import time
    for _ in range(40):
        s = c.get(f"/api/jobs/{j}").json()
        if s["status"] in ("done", "error"):
            break
        time.sleep(.1)
    assert s["status"] == "done" and s["progress"] == 1.0
    j2 = c.post(f"/api/projects/{pid}/run", json={"algorithm": "cvrp", "params": {"capacity": 150, "time_limit": 8}}).json()["job_id"]
    c.post(f"/api/jobs/{j2}/cancel")
    for _ in range(100):
        s = c.get(f"/api/jobs/{j2}").json()
        if s["status"] in ("done", "error", "cancelled"):
            break
        time.sleep(.2)
    assert s["status"] in ("cancelled", "done")


def test_errors_are_clear_spanish(pid):
    r = c.post(f"/api/projects/{pid}/run?wait=true", json={"algorithm": "dijkstra", "params": {"source": "99"}})
    assert r.status_code == 422 and "no existe" in r.json()["error"]
    assert c.post(f"/api/projects/{pid}/run", json={"algorithm": "nope"}).status_code == 422
    assert c.get("/api/projects/xxxx").status_code == 404
    r = c.post("/api/ingest/preview", files=[("files", ("vacio.csv", b"", "text/csv"))])
    assert r.status_code == 422 and "vacío" in json.dumps(r.json(), ensure_ascii=False)


def test_edit_versions_restore(pid):
    p = c.get(f"/api/projects/{pid}").json()
    v0 = p["version"]
    p["nodes"][1]["name"] = "Editado"
    c.put(f"/api/projects/{pid}", json={"nodes": p["nodes"]})
    assert c.get(f"/api/projects/{pid}").json()["nodes"][1]["name"] == "Editado"
    c.post(f"/api/projects/{pid}/restore/{v0}")
    assert c.get(f"/api/projects/{pid}").json()["nodes"][1]["name"] != "Editado"
    assert len(c.get(f"/api/projects/{pid}/versions").json()) >= 3


def test_generate_edges_methods():
    raw = b"id,lat,lon\nA,1,1\nB,1.1,1\nC,1,1.1\nD,1.1,1.1\nE,1.2,1.2\n"
    pv = c.post("/api/ingest/preview", files=[("files", ("n.csv", raw, "text/csv"))]).json()
    pr = c.post("/api/ingest/commit", json={"upload_id": pv["upload_id"], "name": "x"}).json()
    assert any(w["codigo"] == "sin_aristas" for w in pr["meta"]["warnings"])
    for m, extra in (("knn", {"k": 2}), ("delaunay", {}), ("complete", {}), ("threshold", {"threshold": 30})):
        r = c.post(f"/api/projects/{pr['id']}/generate-edges", json={"method": m, **extra}).json()
        assert r["added"] > 0 and r["project"]["edges"], m
    assert len(r["project"]["edges"]) == 10 or m == "threshold"


def test_ask_free_question(pid):
    a = c.post(f"/api/projects/{pid}/ask", json={"question": "¿cuál es la ruta más barata entre Punto 00 y Punto 09 si cierro Punto 03?"}).json()
    assert a["understood"] and "Sin cierre" in a["message"]
    b = c.post(f"/api/projects/{pid}/ask", json={"question": "blablabla"}).json()
    assert b["understood"] is False


def test_all_exports_open(pid):
    run = c.post(f"/api/projects/{pid}/run?wait=true", json={"algorithm": "dijkstra", "params": {"source": "0", "target": "5"}}).json()["run_id"]
    import docx, openpyxl
    for k in ("excel", "word", "pdf", "png", "svg", "graphml", "gexf", "json", "csv", "geojson", "sqlite", "grafos", "zip"):
        r = c.get(f"/api/projects/{pid}/export/{k}", params={"run_id": run})
        assert r.status_code == 200 and len(r.content) > 200, k
        if k == "word":
            d = docx.Document(io.BytesIO(r.content))
            assert any("Resumen ejecutivo" in p.text for p in d.paragraphs) and d.tables
        if k == "excel":
            wb = openpyxl.load_workbook(io.BytesIO(r.content))
            assert {"Nodos", "Aristas", "Parámetros", "Matriz de costos"} <= set(wb.sheetnames)
            assert str(wb["Aristas"]["F2"].value).startswith("=") or wb["Aristas"]["F2"].value is not None
        if k == "pdf":
            assert r.content[:4] == b"%PDF"
        if k == "png":
            assert r.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert c.get(f"/api/projects/{pid}/export/word").status_code == 422


def test_grafos_package_format_and_expected_values():
    p = sample_small()
    files = export.grafos_package(p, scale=100, source="0", target="9")
    mat = files["1_matriz_coste_arcos.txt"]
    assert b"\r\n" in mat and b"\n" not in mat.replace(b"\r\n", b"")
    rows = [r.split("\t") for r in mat.decode().split("\r\n") if r]
    n = len(p["nodes"])
    assert len(rows) == n and all(len(r) == n for r in rows)
    assert rows[0][0] == "#"
    xy = files["3_coordenadas_XY_nodos.txt"].decode().split("\r\n")
    pts = [tuple(map(int, l.split("\t")[1:])) for l in xy if l]
    for i in range(len(pts)):
        for j in range(i + 1, len(pts)):
            assert ((pts[i][0] - pts[j][0]) ** 2 + (pts[i][1] - pts[j][1]) ** 2) ** .5 >= 9.5
    vals = files["4_valores_nodo_2col.txt"].decode().split("\r\n")
    assert all(len(l.split("\t")) == 2 for l in vals if l)
    # valores esperados recalculados desde los enteros exportados
    import networkx as nx
    G = nx.DiGraph()
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if v != "#":
                G.add_edge(i, j, w=int(v))
    tot = sum(d for a in nx.all_pairs_dijkstra_path_length(G, weight="w") for d in a[1].values())
    txt = files["0_INSTRUCCIONES_Grafos.txt"].decode()
    assert f": {tot}" in txt and "SUSTITUIR" in txt and "ACTUALIZAR" in txt and "IZQUIERDO" in txt and "NO VERIFICADO" in txt
    # reimportar lo exportado reproduce el mismo grafo
    from app.ingest import parse_file, build_project
    q = build_project("r", [parse_file("m.txt", mat)])
    assert len(q["edges"]) == len({frozenset((e["source"], e["target"])) for e in p["edges"]})


def test_omml_equation_in_word():
    from app.export import _omml
    from lxml import etree
    x = etree.tostring(_omml("c_{ij} = c_{km}·d_{ij} + t^{2}")).decode()
    assert "sSub" in x and "sSup" in x


def test_templates_download():
    for k in c.get("/api/templates").json():
        r = c.get(f"/api/templates/{k}")
        assert r.status_code == 200 and zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert c.get("/api/templates/zzz").status_code == 404
