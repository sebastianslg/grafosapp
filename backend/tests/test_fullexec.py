import io
import pytest
from fastapi.testclient import TestClient
from app.algos import registry as R
from app.fullexec import build_summary, lower_bound, Builder
from app.main import app
from app.samples import sample_small, sample_fleet, sample_transshipment

REQUIRED = ["Resumen en una mirada", "Conclusión", "Recomendación y siguientes pasos"]


def heads(s):
    return [x["heading"] for x in s["sections"]]


def text(s):
    return " ".join(p for sec in s["sections"] for p in sec["paragraphs"] + sec["bullets"])


CASES = [("sample_fleet", "vrptw", {"time_limit": 2}), ("sample_small", "tsp", {"depot": "0", "time_limit": 2}),
         ("sample_small", "dijkstra", {"source": "0", "target": "9"}), ("sample_small", "kruskal", {}), ("sample_small", "prim", {"mode": "max"}),
         ("sample_small", "max_flow", {"source": "0", "target": "9"}), ("sample_transshipment", "min_cost_flow", {}),
         ("sample_small", "p_median", {"p": 3}), ("sample_small", "articulation", {}), ("sample_fleet", "fstsp", {"time_limit": 2}),
         ("sample_small", "degrees", {}), ("sample_small", "coloring", {}), ("sample_small", "cvrp", {"capacity": 200, "time_limit": 2}),
         ("sample_small", "dijkstra", {"source": "0"})]
SAMPLES = {"sample_fleet": sample_fleet, "sample_small": sample_small, "sample_transshipment": sample_transshipment}


@pytest.mark.parametrize("sname,key,par", CASES)
def test_summary_has_full_structure_for_every_family(sname, key, par):
    p = SAMPLES[sname]()
    r = R.run(key, p, par)
    s = build_summary(p, r, {"light": True})
    h = heads(s)
    assert all(x in h for x in REQUIRED), h
    assert s["sections"][0]["heading"] == "Resumen en una mirada"
    assert any(x.startswith("1. Datos") for x in h) and any("Riesgos" in x for x in h)
    assert s["key_figures"] and all(f["label"] and str(f["value"]) for f in s["key_figures"])
    assert "no son datos operativos oficiales" in text(s) or r["status"] != "heuristic" or key not in ("cvrp", "vrptw")


def test_routing_cost_components_add_up_and_figures_trace_to_result():
    p = sample_fleet()
    r = R.run("vrptw", p, {"time_limit": 2})
    tot = r["value"]
    parts = sum(x["fixed"] + x["cost_dist"] + x["cost_time"] for x in r["paths"])
    assert parts == pytest.approx(tot, rel=1e-6)
    s = build_summary(p, r, {"light": True})
    assert s["key_figures"][0]["value"] == f"{tot:.2f}"
    tbl = next(t for sec in s["sections"] for t in sec["tables"] if t["title"] == "Componentes del costo")
    assert sum(float(row[1]) for row in tbl["rows"]) == pytest.approx(tot, abs=0.02)
    assert sum(float(row[2].replace(" %", "")) for row in tbl["rows"]) == pytest.approx(100, abs=0.3)


def test_full_mode_runs_scenarios_and_reports_noise_and_fleets():
    p = sample_fleet()
    r = R.run("cvrp", p, {"time_limit": 2})
    s = build_summary(p, r, {"time": 2})
    titles = [t["title"] for sec in s["sections"] for t in sec["tables"]]
    assert "Efecto de un aumento de 40 % en la demanda" in titles
    assert any("Ruido de la búsqueda heurística" in t for t in titles)
    assert "Comparación de flotas" in titles
    assert s["scenarios_run"] >= 4 and s["mode"] == "completo"
    noise = next(t for sec in s["sections"] for t in sec["tables"] if t["title"].startswith("Ruido"))
    best = min(float(x[1]) for x in noise["rows"])
    assert all(float(x[1]) >= best for x in noise["rows"])
    fl = next(t for sec in s["sections"] for t in sec["tables"] if t["title"] == "Comparación de flotas")
    assert any(row[0].startswith("Solo ") for row in fl["rows"])


def test_infeasible_growth_scenario_reports_extra_vehicles():
    p = sample_fleet()
    r = R.run("cvrp", p, {"time_limit": 2})
    s = build_summary(p, r, {"time": 2})
    t = text(s)
    assert "no es factible con la flota actual" in t and "vuelve a ser factible" in t


def test_lower_bound_is_valid():
    for p, par in ((sample_small(), {"capacity": 200}), (sample_fleet(), {})):
        r = R.run("cvrp", p, dict(par, time_limit=2))
        b = Builder(p, r, {"light": True})
        lb = lower_bound(b, r["paths"])
        assert lb is not None and lb <= r["value"] + 1e-6


def test_weak_bound_is_not_advertised_as_quality_claim():
    p = sample_fleet()
    r = R.run("vrptw", p, {"time_limit": 2})
    s = build_summary(p, r, {"light": True})
    assert not any("Brecha máxima" in f["label"] for f in s["key_figures"])
    assert "demasiado débil" in text(s)


def test_path_summary_lists_alternatives_and_node_criticality():
    p = sample_small()
    r = R.run("dijkstra", p, {"source": "0", "target": "5"})
    s = build_summary(p, r, {})
    titles = [t["title"] for sec in s["sections"] for t in sec["tables"]]
    assert "Ruta óptima y alternativas" in titles and "Misma consulta con otros criterios" in titles
    assert s["key_figures"][0]["value"] == f"{r['value']:.2f}"


def test_tree_summary_replacement_costs_are_consistent():
    p = sample_small()
    r = R.run("kruskal", p, {})
    s = build_summary(p, r, {})
    t = next(t for sec in s["sections"] for t in sec["tables"] if t["title"].startswith("Costo de reemplazar"))
    for row in t["rows"]:
        if row[4] != "∞":
            assert float(row[4]) >= -1e-9        # reemplazar nunca abarata el árbol mínimo


c = TestClient(app)


def test_summary_endpoints_edit_and_export_roundtrip():
    pid = c.post("/api/samples/pequeno").json()["id"]
    rid = c.post(f"/api/projects/{pid}/run?wait=true", json={"algorithm": "tsp", "params": {"depot": "0", "time_limit": 2}}).json()["run_id"]
    r = c.post(f"/api/runs/{rid}/summary?wait=true", json={"time": 2, "light": True})
    assert r.status_code == 200 and r.json()["summary"]["mode"] == "light"
    s = c.get(f"/api/runs/{rid}").json()["result"]["full_summary"]
    s["sections"][0]["paragraphs"][0] = "TEXTO EDITADO POR EL ESTUDIANTE"
    assert c.put(f"/api/runs/{rid}/summary", json={"summary": s}).status_code == 200
    assert c.get(f"/api/runs/{rid}").json()["result"]["full_summary"]["edited"] is True
    import docx
    d = docx.Document(io.BytesIO(c.get(f"/api/projects/{pid}/export/word", params={"run_id": rid}).content))
    assert any("TEXTO EDITADO POR EL ESTUDIANTE" in p.text for p in d.paragraphs)
    assert c.get(f"/api/projects/{pid}/export/pdf", params={"run_id": rid}).content[:4] == b"%PDF"


def test_summary_job_progress_and_errors():
    pid = c.post("/api/samples/pequeno").json()["id"]
    rid = c.post(f"/api/projects/{pid}/run?wait=true", json={"algorithm": "kruskal", "params": {}}).json()["run_id"]
    jid = c.post(f"/api/runs/{rid}/summary", json={"light": True}).json()["job_id"]
    import time
    for _ in range(50):
        j = c.get(f"/api/jobs/{jid}").json()
        if j["status"] in ("done", "error"):
            break
        time.sleep(.1)
    assert j["status"] == "done" and j["progress"] == 1.0
    assert c.post("/api/runs/inexistente/summary").status_code == 404
    assert c.put(f"/api/runs/{rid}/summary", json={"summary": {"x": 1}}).status_code == 422


def test_word_without_precomputed_summary_still_includes_full_structure():
    pid = c.post("/api/samples/pequeno").json()["id"]
    rid = c.post(f"/api/projects/{pid}/run?wait=true", json={"algorithm": "dijkstra", "params": {"source": "0", "target": "5"}}).json()["run_id"]
    import docx
    d = docx.Document(io.BytesIO(c.get(f"/api/projects/{pid}/export/word", params={"run_id": rid}).content))
    heads_ = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "Resumen ejecutivo" in heads_ and "Conclusión" in heads_ and "Recomendación y siguientes pasos" in heads_ and "Anexo A. Respuesta directa" in heads_
